"""
Test E2E — export Excel / CSV des modules créant des données.

Parcours : bootstrap (admin) → fixtures → login → export des 6 endpoints.

Valide :
- ``GET …/export`` (CSV + XLSX) : statut 200, ``Content-Disposition``
  d'un téléchargement, en‑têtes identiques à ceux de l'import, lignes ;
- ``?template=1`` : modèle vierge (en‑têtes seuls) ;
- ``?format=…`` invalide → 400 ;
- aller‑retour **export → import** : le CSV des contacts se réimporte
  (upsert, 0 création) et celui des commandes recrée 1 commande.

Usage : ``python test_export_flow.py`` (crée puis supprime sa base SQLite).
"""
import asyncio
import csv
import io
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BACKEND = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BACKEND)
os.chdir(BACKEND)

DB_FILE = "test_export_flow.db"
DB_URL = f"sqlite+aiosqlite:///./{DB_FILE}"
# À définir AVANT tout import de app.* (le moteur doit pointer sur la base
# de test) — puis bootstrap de l'admin pour obtenir un vrai token.
os.environ["DATABASE_URL"] = DB_URL
os.environ["BOOTSTRAP_ON_STARTUP"] = "true"
os.environ["FIRST_COMPANY_CODE"] = "EXPO"
os.environ["FIRST_ADMIN_EMAIL"] = "admin@export.test"
os.environ["FIRST_ADMIN_PASSWORD"] = "Admin@12345"

import logging

logging.disable(logging.INFO)  # réduit le bruit des logs SQL pendant le test

failures: list[str] = []


def check(label: str, cond: bool, extra: object = "") -> None:
    status = "OK " if cond else "FAIL"
    suffix = f"  → {extra}" if (not cond and extra) else ""
    print(f"[{status}] {label}{suffix}")
    if not cond:
        failures.append(label)


def parse_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text), delimiter=";"))


# En‑têtes attendus par endpoint (miroir exact des imports / services)
EXPECTED_HEADERS: dict[str, list[str]] = {
    "/crm/partners/export": [
        "nom", "code", "role", "type_entite", "telephone", "email", "site_web",
        "numero_fiscal", "numero_registre", "limite_credit", "delai_paiement",
        "adresse", "ville",
    ],
    "/catalogue/products/export": [
        "reference", "nom", "categorie", "marque", "description", "prix_achat",
        "prix_vente", "prix_ml", "unite", "longueur_std", "vente_au_metre",
        "code_barre", "mode_vente", "type",
    ],
    "/sales/quotes/export": [
        "reference", "client", "date", "statut", "produit", "variante",
        "quantite", "unite", "prix_unitaire", "remise", "sous_total",
        "montant_document", "devise", "description",
    ],
    "/sales/orders/export": [
        "reference", "client", "date", "statut", "produit", "variante",
        "quantite", "unite", "prix_unitaire", "remise", "sous_total",
        "montant_document", "devise", "description",
    ],
    "/invoicing/invoices/export": [
        "reference", "client", "date", "echeance", "statut", "produit",
        "variante", "quantite", "unite", "prix_unitaire", "remise",
        "sous_total", "montant_ht", "tva", "montant_document", "montant_paye",
        "devise", "description",
    ],
    "/delivery/notes/export": [
        "reference", "client", "commande", "date", "statut", "produit",
        "variante", "quantite", "quantite_commandee", "unite", "description",
    ],
}


async def main() -> int:
    from httpx import ASGITransport, AsyncClient
    from sqlalchemy import select

    from app.core.database import AsyncSessionLocal, Base, engine
    from app.core.services.seed_service import run_bootstrap_if_needed
    from app.main import app as fastapi_app

    import app.core.models  # noqa: F401
    import app.modules.catalogue.models  # noqa: F401
    import app.modules.crm.models  # noqa: F401
    import app.modules.sales.models  # noqa: F401
    import app.modules.stock.models  # noqa: F401
    import app.modules.invoicing.models  # noqa: F401
    import app.modules.delivery.models  # noqa: F401

    db_path = os.path.join(BACKEND, DB_FILE)
    if os.path.exists(db_path):
        os.remove(db_path)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    # ------------------------------------------------------------------ #
    # 1. Bootstrap (entreprise + admin) puis fixtures dans SON entreprise
    #    (le filtre multi‑tenant automatique masque les autres sociétés)
    # ------------------------------------------------------------------ #
    print("Bootstrap :", await run_bootstrap_if_needed())

    async with AsyncSessionLocal() as session:
        from app.core.models.company import Company
        from app.core.models.user import User
        from app.modules.catalogue.models.product import (
            Product,
            ProductCategory,
            ProductUnit,
        )
        from app.modules.crm.models.partner import Partner, PartnerAddress
        from app.modules.delivery.services.delivery_service import DeliveryService
        from app.modules.invoicing.services.invoice_service import InvoiceService
        from app.modules.sales.schemas import (
            OrderLineCreate,
            QuoteCreate,
            QuoteLineCreate,
            SalesOrderCreate,
        )
        from app.modules.sales.services.order_service import OrderService
        from app.modules.sales.services.quote_service import QuoteService

        company = (await session.execute(
            select(Company).order_by(Company.id)
        )).scalars().first()
        admin = (await session.execute(
            select(User).where(User.email == "admin@export.test")
        )).scalars().first()
        check("Fixtures : entreprise bootstrap récupérée", company is not None)
        check("Fixtures : admin récupéré", admin is not None)

        partner = Partner(
            name="Client Export", code="EXP-C1", role="client",
            entity_type="company", phone="+225 01 02 03 04",
            email="client@export.test", company_id=company.id,
        )
        session.add(partner)
        await session.flush()
        session.add(PartnerAddress(
            partner_id=partner.id, line_1="Boulevard Export 12",
            city="Abidjan", is_primary=True, company_id=company.id,
        ))

        category = ProductCategory(
            name="Catégorie Export", sale_mode="unit", company_id=company.id,
        )
        unit = ProductUnit(name="Pièce", code="pz", company_id=company.id)
        session.add_all([category, unit])
        await session.flush()

        product = Product(
            reference="EXP-001", name="Produit Export", sale_mode="unit",
            selling_price=1000, category_id=category.id,
            sale_unit_id=unit.id, company_id=company.id,
        )
        session.add(product)
        await session.flush()

        quote = await QuoteService(session).create_quote(QuoteCreate(
            partner_id=partner.id,
            lines=[QuoteLineCreate(product_id=product.id, quantity=2,
                                   unit_price=1000)],
        ), admin)
        order = await OrderService(session).create_order(SalesOrderCreate(
            partner_id=partner.id,
            lines=[OrderLineCreate(product_id=product.id, quantity_ordered=3,
                                   unit_price=1000)],
        ))
        invoice = await InvoiceService(session).create_from_order(order.id)
        delivery = await DeliveryService(session).create_from_order(order.id)
        check("Fixtures : devis / commande / facture / BL créés",
              all(x is not None for x in (quote, order, invoice, delivery)))

    # ------------------------------------------------------------------ #
    # 2. Login + exports (CSV, XLSX, modèle vierge, format invalide)
    # ------------------------------------------------------------------ #
    transport = ASGITransport(app=fastapi_app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        login = await client.post(
            "/api/v1/core/auth/login",
            json={"email": "admin@export.test", "password": "Admin@12345"},
        )
        check("Login admin", login.status_code == 200, login.text[:300])
        if login.status_code != 200:
            return 1
        auth = {"Authorization": f"Bearer {login.json()['access_token']}"}
        bodies: dict[str, bytes] = {}

        for path, headers in EXPECTED_HEADERS.items():
            url = f"/api/v1{path}"

            # --- CSV ---
            res = await client.get(f"{url}?format=csv", headers=auth)
            check(f"{path} → CSV 200", res.status_code == 200,
                  res.status_code != 200 and res.text[:300])
            if res.status_code != 200:
                continue
            disposition = res.headers.get("content-disposition", "")
            check(f"{path} → téléchargement nommé (.csv)",
                  "attachment" in disposition and ".csv" in disposition,
                  disposition)
            check(f"{path} → type text/csv",
                  "text/csv" in res.headers.get("content-type", ""),
                  res.headers.get("content-type"))
            rows = parse_csv(res.content)
            check(f"{path} → en‑têtes conformes", bool(rows) and rows[0] == headers,
                  rows[:1])
            check(f"{path} → au moins 1 ligne de données", len(rows) >= 2, len(rows))
            bodies[path] = res.content

            # --- XLSX ---
            res_x = await client.get(f"{url}?format=xlsx", headers=auth)
            check(f"{path} → XLSX 200", res_x.status_code == 200,
                  res_x.status_code != 200 and res_x.text[:300])
            if res_x.status_code == 200:
                from openpyxl import load_workbook

                wb = load_workbook(io.BytesIO(res_x.content), read_only=True)
                sheet = wb.worksheets[0]
                xlsx_rows = list(sheet.iter_rows(values_only=True))
                wb.close()
                check(f"{path} → en‑têtes XLSX conformes",
                      list(xlsx_rows[0]) == headers, xlsx_rows[:1])
                check(f"{path} → au moins 1 ligne XLSX",
                      len(xlsx_rows) >= 2, len(xlsx_rows))
                check(f"{path} → disposition .xlsx",
                      ".xlsx" in res_x.headers.get("content-disposition", ""),
                      res_x.headers.get("content-disposition"))

            # --- Modèle vierge (en‑têtes seuls) ---
            res_t = await client.get(f"{url}?format=csv&template=1", headers=auth)
            check(f"{path} → modèle vierge 200", res_t.status_code == 200)
            if res_t.status_code == 200:
                template_rows = parse_csv(res_t.content)
                check(f"{path} → modèle = en‑têtes seuls",
                      template_rows == [headers], template_rows)

        # Format inconnu → 400
        res_bad = await client.get(
            "/api/v1/crm/partners/export?format=pdf", headers=auth
        )
        check("Format inconnu → 400", res_bad.status_code == 400,
              res_bad.status_code)

        # Filtre d'écran repris (onglets Clients / Fournisseurs)
        res_client = await client.get(
            "/api/v1/crm/partners/export?format=csv&role=client", headers=auth
        )
        check("Filtre role=client → lignes exportées",
              res_client.status_code == 200
              and len(parse_csv(res_client.content)) >= 2,
              res_client.status_code)
        res_supplier = await client.get(
            "/api/v1/crm/partners/export?format=csv&role=supplier", headers=auth
        )
        check("Filtre role=supplier → aucun fournisseur",
              res_supplier.status_code == 200
              and len(parse_csv(res_supplier.content)) == 1,
              parse_csv(res_supplier.content))

        # ------------------------------------------------------------------ #
        # 3. Aller‑retour : le fichier exporté se réimporte tel quel
        # ------------------------------------------------------------------ #
        res_imp = await client.post(
            "/api/v1/crm/partners/import",
            files={"file": ("contacts.csv", bodies["/crm/partners/export"],
                            "text/csv")},
            headers=auth,
        )
        check("Aller‑retour contacts : import 200", res_imp.status_code == 200,
              res_imp.text[:300])
        if res_imp.status_code == 200:
            data = res_imp.json()
            check("Aller‑retour contacts : upsert sans création ni erreur",
                  data.get("created") == 0
                  and (data.get("updated") or 0) >= 1
                  and not data.get("errors"),
                  data)

        res_ord = await client.post(
            "/api/v1/sales/orders/import",
            files={"file": ("commandes.csv", bodies["/sales/orders/export"],
                            "text/csv")},
            headers=auth,
        )
        check("Aller‑retour commandes : import 200", res_ord.status_code == 200,
              res_ord.text[:300])
        if res_ord.status_code == 200:
            data = res_ord.json()
            check("Aller‑retour commandes : 1 commande recréée, 0 erreur",
                  data.get("orders_created") == 1 and not data.get("errors"),
                  data)

    await engine.dispose()
    if os.path.exists(db_path):
        os.remove(db_path)

    print("\n" + "=" * 60)
    if failures:
        print(f"ÉCHEC — {len(failures)} contrôle(s) en échec :")
        for item in failures:
            print(f"  - {item}")
    else:
        print("SUCCÈS — tous les contrôles sont passés")
    print("=" * 60)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
