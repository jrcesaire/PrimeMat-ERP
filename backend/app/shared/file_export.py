"""
Génération de fichiers d'export (Excel / CSV) — miroir de ``file_import``.

L'objectif est le **transfert de données** : un export doit pouvoir être
réimporté tel quel. Les conventions sont donc identiques à la lecture :

- en‑têtes en français, redescendus au même format normalisé à la
  réimportation (« Prix d'achat » → ``prix_d_achat``) ;
- dates au format ISO ``AAAA-MM-JJ`` (accepté par ``to_date``) ;
- booléens en « oui » / « non » (acceptés par ``to_bool``) ;
- nombres sans zéro parasite (``1500`` et non ``1500.0000``).

Formats :
- ``.csv`` : séparateur ``;`` (Excel francophone) + BOM UTF‑8, lisible par
  Excel comme par l'import ;
- ``.xlsx`` : openpyxl (en‑tête mis en forme, volette figée, filtre),
  repli CSV automatique si la librairie n'est pas installée.
"""
from __future__ import annotations

import codecs
import csv
import io
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from typing import Any, Sequence

from fastapi import HTTPException, Response

from app.core.documents.naming import content_disposition

EXPORT_FORMATS = ("xlsx", "csv")
DEFAULT_FORMAT = "xlsx"

_MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def normalize_format(fmt: str | None) -> str:
    """``"XLSX"`` / ``".csv"`` → ``"xlsx"`` / ``"csv"`` (ou 400 si inconnu)."""
    value = (fmt or DEFAULT_FORMAT).strip().lower().lstrip(".")
    if value not in EXPORT_FORMATS:
        raise HTTPException(
            400, "Format d'export non supporté — utilisez « xlsx » ou « csv »"
        )
    return value


def to_csv_cell(value: Any) -> str:
    """Cellule destinée au CSV (chaîne, séparateur ``;``)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, Enum):
        return str(value.value)
    if isinstance(value, Decimal):
        return _number_str(float(value))
    if isinstance(value, float):
        return _number_str(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def to_xlsx_cell(value: Any) -> Any:
    """Cellule destinée à Excel (types natifs, dates en ISO comme en CSV)."""
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, Enum):
        return str(value.value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _number_str(value: float) -> str:
    """``1500.0 → "1500"``, ``0.5 → "0.5"`` (CSV : reste lisible/parsable)."""
    text = f"{value:.4f}".rstrip("0").rstrip(".")
    return text if text not in ("", "-") else "0"


def export_response(
    fmt: str | None,
    filename_base: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> Response:
    """Réponse de téléchargement (``Content-Disposition: attachment``).

    ``filename_base`` : nom sans extension (« Contacts ») → le fichier est
    daté du jour (« Contacts_2026-10-02.xlsx ») pour classer les exports.
    """
    export_format = normalize_format(fmt)
    filename = f"{filename_base}_{date.today().isoformat()}"

    if export_format == "xlsx":
        try:
            import openpyxl  # noqa: F401
        except ImportError:  # pragma: no cover — repli si librairie absente
            export_format = "csv"
        else:
            content = _xlsx_bytes(headers, rows)

    if export_format == "csv":
        content = _csv_bytes(headers, rows)

    filename = f"{filename}.{export_format}"
    return Response(
        content=content,
        media_type=_MEDIA_TYPES[export_format],
        headers={"Content-Disposition": content_disposition(filename, inline=False)},
    )


def _csv_bytes(
    headers: Sequence[str], rows: Sequence[Sequence[Any]]
) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(
        buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n"
    )
    writer.writerow(list(headers))
    for row in rows:
        writer.writerow([to_csv_cell(cell) for cell in row])
    # BOM : Excel (Windows) détecte alors l'UTF‑8 au lieu de casser les accents
    return (codecs.BOM_UTF8 + buffer.getvalue().encode("utf-8"))


def _xlsx_bytes(
    headers: Sequence[str], rows: Sequence[Sequence[Any]]
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    sheet.append(list(headers))

    header_fill = PatternFill("solid", fgColor="0A2A6B")
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        sheet.append([to_xlsx_cell(cell) for cell in row])

    sheet.freeze_panes = "A2"
    if headers:
        last = get_column_letter(len(headers))
        sheet.auto_filter.ref = f"A1:{last}{max(len(rows) + 1, 1)}"

    # Largeurs : contenu de la colonne, plafonnées pour rester lisibles
    for index, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows[:500]:
            if index - 1 < len(row):
                longest = max(longest, len(str(to_csv_cell(row[index - 1]))))
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(longest + 2, 10), 48
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
